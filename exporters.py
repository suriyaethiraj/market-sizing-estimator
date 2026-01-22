from __future__ import annotations
import io
import json
from schema import MarketSizingReport

NAVY   = "#1F3C6B"
INDIGO = "#4F46E5"
LIGHT  = "#EEF2FF"
GREEN  = "#16A34A"
AMBER  = "#D97706"
RED    = "#DC2626"
MUTED  = "#64748B"
WHITE  = "#FFFFFF"


def _fmt(val: float) -> str:
    """Format USD millions to human-readable."""
    if val >= 1_000_000:
        return f"${val/1_000_000:.1f}T"
    if val >= 1_000:
        return f"${val/1_000:.1f}B"
    return f"${val:.0f}M"


# ── JSON ─────────────────────────────────────────────────────────────────────

def export_json(report: MarketSizingReport) -> str:
    return report.model_dump_json(indent=2)


# ── PDF ──────────────────────────────────────────────────────────────────────

def export_pdf(report: MarketSizingReport) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
        HRFlowable, PageBreak
    )
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=18*mm, rightMargin=18*mm,
                            topMargin=18*mm, bottomMargin=18*mm)

    C_NAVY   = colors.HexColor(NAVY)
    C_INDIGO = colors.HexColor(INDIGO)
    C_LIGHT  = colors.HexColor(LIGHT)
    C_GREEN  = colors.HexColor(GREEN)
    C_MUTED  = colors.HexColor(MUTED)

    title_sty  = ParagraphStyle("T",  fontSize=20, textColor=C_NAVY,   fontName="Helvetica-Bold", spaceAfter=4,  alignment=TA_CENTER)
    sub_sty    = ParagraphStyle("S",  fontSize=11, textColor=C_MUTED,  spaceAfter=12, alignment=TA_CENTER)
    h2_sty     = ParagraphStyle("H2", fontSize=13, textColor=C_NAVY,   fontName="Helvetica-Bold", spaceBefore=12, spaceAfter=6)
    h3_sty     = ParagraphStyle("H3", fontSize=10, textColor=C_INDIGO, fontName="Helvetica-Bold", spaceBefore=8,  spaceAfter=4)
    body_sty   = ParagraphStyle("B",  fontSize=9,  textColor=colors.HexColor("#1E293B"), spaceAfter=4, leading=14, alignment=TA_JUSTIFY)
    small_sty  = ParagraphStyle("Sm", fontSize=8,  textColor=C_MUTED,  spaceAfter=3)

    def hdr_tbl(text):
        t = Table([[Paragraph(text, ParagraphStyle("HT", fontSize=12, textColor=colors.white,
                                                    fontName="Helvetica-Bold"))]], colWidths=[174*mm])
        t.setStyle(TableStyle([
            ("BACKGROUND",    (0,0),(-1,-1), C_NAVY),
            ("TOPPADDING",    (0,0),(-1,-1), 8),
            ("BOTTOMPADDING", (0,0),(-1,-1), 8),
            ("LEFTPADDING",   (0,0),(-1,-1), 10),
        ]))
        return t

    def scenario_table(label, tam, sam, som):
        data = [
            ["",     "Low",              "Mid (Base)",        "High"],
            ["TAM",  _fmt(tam.low),      _fmt(tam.mid),       _fmt(tam.high)],
            ["SAM",  _fmt(sam.low),      _fmt(sam.mid),       _fmt(sam.high)],
            ["SOM",  _fmt(som.low),      _fmt(som.mid),       _fmt(som.high)],
        ]
        t = Table(data, colWidths=[25*mm, 45*mm, 55*mm, 45*mm])
        t.setStyle(TableStyle([
            ("BACKGROUND",    (0,0),(-1,0),  C_NAVY),
            ("TEXTCOLOR",     (0,0),(-1,0),  colors.white),
            ("FONTNAME",      (0,0),(-1,0),  "Helvetica-Bold"),
            ("BACKGROUND",    (0,1),(0,-1),  C_LIGHT),
            ("FONTNAME",      (0,1),(0,-1),  "Helvetica-Bold"),
            ("TEXTCOLOR",     (0,1),(0,-1),  C_NAVY),
            ("BACKGROUND",    (2,1),(2,-1),  colors.HexColor("#F0FDF4")),
            ("FONTNAME",      (2,1),(2,-1),  "Helvetica-Bold"),
            ("FONTSIZE",      (0,0),(-1,-1), 9),
            ("GRID",          (0,0),(-1,-1), 0.5, colors.HexColor("#CBD5E1")),
            ("TOPPADDING",    (0,0),(-1,-1), 5),
            ("BOTTOMPADDING", (0,0),(-1,-1), 5),
            ("LEFTPADDING",   (0,0),(-1,-1), 6),
            ("ALIGN",         (1,0),(-1,-1), "CENTER"),
        ]))
        return t

    story = []

    # Cover
    story += [
        Spacer(1, 16*mm),
        Paragraph(report.title, title_sty),
        Paragraph(f"Industry: {report.industry}  ·  Geographies: {', '.join(report.geographies)}", sub_sty),
        Paragraph(f"Generated: {report.generated_at}  ·  Confidence: {report.confidence_rating}", sub_sty),
        HRFlowable(width="100%", thickness=2, color=C_INDIGO),
        Spacer(1, 6*mm),
    ]

    # Investor narrative box
    story.append(hdr_tbl("💼 Investor Narrative"))
    story.append(Paragraph(report.investor_narrative, ParagraphStyle("IN", fontSize=10,
        textColor=colors.HexColor("#1E293B"), spaceAfter=8, spaceBefore=6, leading=16,
        backColor=colors.HexColor("#EEF2FF"), borderPadding=8)))

    # Executive summary
    story += [Paragraph("Executive Summary", h2_sty), Paragraph(report.executive_summary, body_sty)]

    # Reconciled results
    story += [Spacer(1,4*mm), hdr_tbl("📊 Reconciled Market Size (All Methodologies)"), Spacer(1,3*mm)]
    story.append(scenario_table("Reconciled",
                                report.reconciled_tam, report.reconciled_sam, report.reconciled_som))

    story.append(PageBreak())

    # Per-methodology
    story.append(Paragraph("Methodology Breakdown", h2_sty))
    for mr in report.methodology_results:
        story += [Paragraph(f"{mr.methodology} Approach", h3_sty),
                  Paragraph(mr.narrative, body_sty)]
        story.append(scenario_table(mr.methodology, mr.tam, mr.sam, mr.som))
        story.append(Spacer(1, 2*mm))

        if mr.key_assumptions:
            story.append(Paragraph("Key Assumptions:", ParagraphStyle("KA", fontSize=9,
                fontName="Helvetica-Bold", textColor=C_NAVY, spaceAfter=3, spaceBefore=4)))
            assump_data = [["Assumption", "Value", "Source", "Confidence"]]
            for a in mr.key_assumptions:
                assump_data.append([a.label, a.value, a.source, a.confidence])
            at = Table(assump_data, colWidths=[55*mm, 40*mm, 55*mm, 22*mm])
            at.setStyle(TableStyle([
                ("BACKGROUND",    (0,0),(-1,0),  C_LIGHT),
                ("FONTNAME",      (0,0),(-1,0),  "Helvetica-Bold"),
                ("FONTSIZE",      (0,0),(-1,-1), 8),
                ("GRID",          (0,0),(-1,-1), 0.5, colors.HexColor("#CBD5E1")),
                ("TOPPADDING",    (0,0),(-1,-1), 4),
                ("BOTTOMPADDING", (0,0),(-1,-1), 4),
                ("LEFTPADDING",   (0,0),(-1,-1), 5),
                ("VALIGN",        (0,0),(-1,-1), "TOP"),
            ]))
            story.append(at)
        story.append(Spacer(1, 4*mm))

    story.append(PageBreak())

    # 5-year projection
    story += [hdr_tbl("📈 5-Year Market Projection"), Spacer(1, 3*mm)]
    proj_data = [["Year", "TAM", "SAM", "SOM", "CAGR Applied"]]
    for p in report.five_year_projection:
        proj_data.append([str(p.year), _fmt(p.tam), _fmt(p.sam), _fmt(p.som), f"{p.cagr_applied:.1f}%"])
    pt = Table(proj_data, colWidths=[25*mm, 40*mm, 40*mm, 40*mm, 30*mm])
    pt.setStyle(TableStyle([
        ("BACKGROUND",    (0,0),(-1,0), C_NAVY),
        ("TEXTCOLOR",     (0,0),(-1,0), colors.white),
        ("FONTNAME",      (0,0),(-1,0), "Helvetica-Bold"),
        ("FONTSIZE",      (0,0),(-1,-1), 9),
        ("GRID",          (0,0),(-1,-1), 0.5, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS",(0,1),(-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("ALIGN",         (1,0),(-1,-1), "CENTER"),
        ("TOPPADDING",    (0,0),(-1,-1), 5),
        ("BOTTOMPADDING", (0,0),(-1,-1), 5),
    ]))
    story.append(pt)

    # Key risks
    story += [Spacer(1,6*mm), Paragraph("Key Risks", h2_sty)]
    for risk in report.key_risks:
        story.append(Paragraph(f"• {risk}", body_sty))

    # Data sources
    story += [Spacer(1,4*mm), Paragraph("Data Sources Used", h2_sty)]
    for src in report.world_bank_data_used[:15]:
        story.append(Paragraph(f"• {src}", small_sty))

    doc.build(story)
    return buf.getvalue()


# ── EXCEL ─────────────────────────────────────────────────────────────────────

def export_excel(report: MarketSizingReport) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb  = openpyxl.Workbook()
    thin = Side(style="thin", color="CBD5E1")
    bdr  = Border(left=thin, right=thin, top=thin, bottom=thin)

    def nfill(hex_): return PatternFill("solid", fgColor=hex_.lstrip("#"))
    def nfont(color="FFFFFF", bold=False, size=10):
        return Font(color=color.lstrip("#"), bold=bold, size=size)
    def acell(ws, row, col, val, fill=None, font=None, align="left", wrap=False):
        c = ws.cell(row=row, column=col, value=val)
        c.border = bdr
        c.alignment = Alignment(horizontal=align, vertical="top", wrap_text=wrap)
        if fill: c.fill = fill
        if font: c.font = font
        return c

    # ── Sheet 1: Summary ─────────────────────────────────────────────────────
    ws = wb.active; ws.title = "Summary"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 20

    r = 1
    ws.merge_cells(f"A{r}:D{r}")
    acell(ws, r, 1, report.title, nfill(NAVY), nfont(bold=True, size=13), "center")
    r += 1
    acell(ws, r, 1, f"Industry: {report.industry}", font=nfont("64748B", size=9))
    acell(ws, r, 2, f"Geographies: {', '.join(report.geographies)}", font=nfont("64748B", size=9))
    acell(ws, r, 3, f"Generated: {report.generated_at}", font=nfont("64748B", size=9))
    acell(ws, r, 4, f"Confidence: {report.confidence_rating}", font=nfont("64748B", size=9))
    r += 2

    # Investor narrative
    ws.merge_cells(f"A{r}:D{r}")
    acell(ws, r, 1, "INVESTOR NARRATIVE", nfill(INDIGO), nfont(bold=True), "center"); r+=1
    ws.merge_cells(f"A{r}:D{r}")
    ws.row_dimensions[r].height = 60
    acell(ws, r, 1, report.investor_narrative, nfill(LIGHT), font=Font(size=9), wrap=True); r+=2

    # Reconciled table
    ws.merge_cells(f"A{r}:D{r}")
    acell(ws, r, 1, "RECONCILED MARKET SIZE", nfill(NAVY), nfont(bold=True), "center"); r+=1
    for hdr in ["Metric","Low","Mid (Base)","High"]:
        acell(ws, r, ["Metric","Low","Mid (Base)","High"].index(hdr)+1, hdr,
              nfill(LIGHT), nfont("1F3C6B", bold=True), "center")
    r+=1
    for label, scenario in [("TAM",report.reconciled_tam),("SAM",report.reconciled_sam),("SOM",report.reconciled_som)]:
        acell(ws, r, 1, label, font=Font(bold=True, color="1F3C6B"))
        acell(ws, r, 2, _fmt(scenario.low),  align="center")
        acell(ws, r, 3, _fmt(scenario.mid),  align="center",
              fill=nfill("F0FDF4") if label=="SOM" else None,
              font=Font(bold=True, color="16A34A") if label=="SOM" else None)
        acell(ws, r, 4, _fmt(scenario.high), align="center")
        r+=1

    r+=1
    acell(ws, r, 1, "Executive Summary", nfill(LIGHT), nfont("1F3C6B", bold=True)); r+=1
    ws.merge_cells(f"A{r}:D{r}")
    ws.row_dimensions[r].height = 72
    acell(ws, r, 1, report.executive_summary, wrap=True, font=Font(size=9)); r+=2

    acell(ws, r, 1, "Key Risks", nfill(LIGHT), nfont("1F3C6B", bold=True)); r+=1
    for risk in report.key_risks:
        ws.merge_cells(f"A{r}:D{r}")
        acell(ws, r, 1, f"• {risk}", font=Font(size=9)); r+=1

    # ── Sheet 2: Methodologies ────────────────────────────────────────────────
    ws2 = wb.create_sheet("Methodologies")
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 15
    ws2.column_dimensions["C"].width = 15
    ws2.column_dimensions["D"].width = 15
    ws2.column_dimensions["E"].width = 30
    ws2.column_dimensions["F"].width = 20
    ws2.column_dimensions["G"].width = 20
    ws2.column_dimensions["H"].width = 12

    r2 = 1
    for mr in report.methodology_results:
        ws2.merge_cells(f"A{r2}:H{r2}")
        acell(ws2, r2, 1, f"{mr.methodology} Approach", nfill(NAVY), nfont(bold=True), "center"); r2+=1
        for hdr,ci in [("",1),("Low",2),("Mid",3),("High",4)]:
            acell(ws2, r2, ci, hdr, nfill(LIGHT), nfont("1F3C6B", bold=True), "center")
        r2+=1
        for lbl, scen in [("TAM",mr.tam),("SAM",mr.sam),("SOM",mr.som)]:
            acell(ws2, r2, 1, lbl, font=Font(bold=True, color="1F3C6B"))
            for ci, v in enumerate([scen.low, scen.mid, scen.high], 2):
                acell(ws2, r2, ci, _fmt(v), align="center"); r2+=1

        ws2.merge_cells(f"A{r2}:H{r2}")
        ws2.row_dimensions[r2].height = 48
        acell(ws2, r2, 1, mr.narrative, wrap=True, font=Font(size=9)); r2+=1

        if mr.key_assumptions:
            for hdr,ci in [("Assumption",1),("Value",2),("Source",3),("Confidence",4)]:
                acell(ws2, r2, ci, hdr, nfill(LIGHT), nfont("1F3C6B", bold=True))
            r2+=1
            for a in mr.key_assumptions:
                acell(ws2, r2, 1, a.label); acell(ws2, r2, 2, a.value)
                acell(ws2, r2, 3, a.source); acell(ws2, r2, 4, a.confidence); r2+=1
        r2+=1

    # ── Sheet 3: 5-Year Projection ────────────────────────────────────────────
    ws3 = wb.create_sheet("5-Year Projection")
    for col,w in zip("ABCDE",[10,18,18,18,15]):
        ws3.column_dimensions[col].width = w
    for ci,hdr in enumerate(["Year","TAM","SAM","SOM","CAGR"],1):
        acell(ws3, 1, ci, hdr, nfill(NAVY), nfont(bold=True), "center")
    for ri, p in enumerate(report.five_year_projection, 2):
        acell(ws3, ri, 1, p.year, align="center")
        acell(ws3, ri, 2, _fmt(p.tam), align="center")
        acell(ws3, ri, 3, _fmt(p.sam), align="center")
        acell(ws3, ri, 4, _fmt(p.som), align="center",
              fill=nfill("F0FDF4"), font=Font(bold=True, color="16A34A"))
        acell(ws3, ri, 5, f"{p.cagr_applied:.1f}%", align="center")

    # ── Sheet 4: Sensitivity ──────────────────────────────────────────────────
    ws4 = wb.create_sheet("Sensitivity Analysis")
    r4 = 1
    for axis in report.sensitivity_axes:
        ws4.merge_cells(f"A{r4}:F{r4}")
        acell(ws4, r4, 1, f"Sensitivity: {axis.variable}", nfill(NAVY), nfont(bold=True)); r4+=1
        for ci, (val, impact) in enumerate(zip(axis.values, axis.som_impacts), 1):
            acell(ws4, r4, ci, val,          nfill(LIGHT), nfont("1F3C6B", bold=True), "center")
            acell(ws4, r4+1, ci, _fmt(impact), align="center")
        r4+=3

    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


# ── POWERPOINT ────────────────────────────────────────────────────────────────

def export_pptx(report: MarketSizingReport) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    def rgb(hex_): h=hex_.lstrip("#"); return RGBColor(int(h[0:2],16),int(h[2:4],16),int(h[4:6],16))
    def add_text(tf, text, size=18, bold=False, color=NAVY, align=PP_ALIGN.LEFT):
        p = tf.add_paragraph(); p.text = text; p.alignment = align
        r = p.runs[0] if p.runs else p.add_run()
        r.text = text; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(color)
        return p

    prs = Presentation(); prs.slide_width = Inches(13.33); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def new_slide(): return prs.slides.add_slide(blank)

    def bg(slide, color=NAVY):
        from pptx.util import Emu
        bg_ = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
        bg_.fill.solid(); bg_.fill.fore_color.rgb = rgb(color)
        bg_.line.fill.background()

    def box(slide, l,t,w,h,text,bg_hex=WHITE,fg_hex=NAVY,size=14,bold=False,align=PP_ALIGN.LEFT):
        shp = slide.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
        shp.fill.solid(); shp.fill.fore_color.rgb = rgb(bg_hex)
        tf = shp.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = text; p.alignment = align
        r = p.runs[0] if p.runs else p.add_run()
        r.text = text; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(fg_hex)

    # Slide 1 — Title
    s1 = new_slide(); bg(s1, NAVY)
    box(s1,0.5,1.2,12,1.2, report.title,          NAVY,   WHITE,  32, True, PP_ALIGN.CENTER)
    box(s1,0.5,2.6,12,0.6, f"{report.industry}  ·  {', '.join(report.geographies)}", NAVY, LIGHT, 16, False, PP_ALIGN.CENTER)
    box(s1,0.5,3.4,12,0.5, f"Confidence: {report.confidence_rating}  ·  {report.generated_at}", NAVY, MUTED, 12, False, PP_ALIGN.CENTER)
    box(s1,1.0,4.2,11,1.8, report.investor_narrative, INDIGO, WHITE, 15, False, PP_ALIGN.CENTER)

    # Slide 2 — Reconciled TAM/SAM/SOM
    s2 = new_slide(); bg(s2, WHITE)
    box(s2,0.3,0.2,12,0.7,"📊 Reconciled Market Size", LIGHT, NAVY, 22, True)
    metrics = [
        ("TAM", report.reconciled_tam),
        ("SAM", report.reconciled_sam),
        ("SOM", report.reconciled_som),
    ]
    for i,(lbl,sc) in enumerate(metrics):
        x = 0.5 + i*4.2
        box(s2, x, 1.1, 3.8, 0.6, lbl, NAVY, WHITE, 16, True, PP_ALIGN.CENTER)
        box(s2, x, 1.8, 3.8, 0.9, _fmt(sc.mid), INDIGO if lbl=="SOM" else LIGHT,
            WHITE if lbl=="SOM" else NAVY, 28, True, PP_ALIGN.CENTER)
        box(s2, x, 2.8, 3.8, 0.5, f"Low: {_fmt(sc.low)}  ·  High: {_fmt(sc.high)}", "#F8FAFC", MUTED, 11, False, PP_ALIGN.CENTER)

    box(s2,0.3,4.0,12,2.8, report.executive_summary, "#F8FAFC", "#1E293B", 12, False)

    # Slide 3 — Methodology comparison
    s3 = new_slide(); bg(s3, WHITE)
    box(s3,0.3,0.2,12,0.7,"🔬 Methodology Comparison", LIGHT, NAVY, 22, True)
    for i,mr in enumerate(report.methodology_results):
        x = 0.4 + i*4.2
        box(s3, x, 1.1, 3.8, 0.6, mr.methodology, NAVY, WHITE, 13, True, PP_ALIGN.CENTER)
        box(s3, x, 1.8, 3.8, 0.5, f"TAM: {_fmt(mr.tam.mid)}", LIGHT, NAVY, 11, False, PP_ALIGN.CENTER)
        box(s3, x, 2.4, 3.8, 0.5, f"SAM: {_fmt(mr.sam.mid)}", LIGHT, NAVY, 11, False, PP_ALIGN.CENTER)
        box(s3, x, 3.0, 3.8, 0.5, f"SOM: {_fmt(mr.som.mid)}", "#F0FDF4", GREEN, 12, True, PP_ALIGN.CENTER)
        box(s3, x, 3.7, 3.8, 2.2, mr.narrative, "#F8FAFC", MUTED, 9, False)

    # Slide 4 — 5-Year Projection
    s4 = new_slide(); bg(s4, WHITE)
    box(s4,0.3,0.2,12,0.7,"📈 5-Year Market Projection", LIGHT, NAVY, 22, True)
    hdrs = ["Year","TAM","SAM","SOM","CAGR"]
    for ci,h in enumerate(hdrs):
        box(s4, 0.3+ci*2.4, 1.1, 2.2, 0.5, h, NAVY, WHITE, 11, True, PP_ALIGN.CENTER)
    for ri,p in enumerate(report.five_year_projection):
        vals = [str(p.year), _fmt(p.tam), _fmt(p.sam), _fmt(p.som), f"{p.cagr_applied:.1f}%"]
        bg_c = LIGHT if ri%2==0 else WHITE
        for ci,v in enumerate(vals):
            box(s4, 0.3+ci*2.4, 1.7+ri*0.9, 2.2, 0.8, v,
                "#F0FDF4" if ci==3 else bg_c,
                GREEN if ci==3 else NAVY, 11, ci==3, PP_ALIGN.CENTER)

    # Slide 5 — Key Risks
    s5 = new_slide(); bg(s5, WHITE)
    box(s5,0.3,0.2,12,0.7,"⚠️ Key Risks & Considerations", LIGHT, NAVY, 22, True)
    risk_text = "\n".join([f"• {r}" for r in report.key_risks])
    box(s5,0.5,1.2,12,5.0, risk_text, "#FFF7ED", "#1E293B", 13, False)

    buf = io.BytesIO(); prs.save(buf); return buf.getvalue()
